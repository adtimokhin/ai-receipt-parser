"""Non-AI 529 report builder: the .xlsx that breaks rent out from food."""

from __future__ import annotations

import io
from datetime import UTC, datetime

from openpyxl import load_workbook  # type: ignore[import-untyped]

from receipt_parser_backend.receipts.models import Receipt, ReceiptFiles
from receipt_parser_backend.reports.xlsx_builder import build_report_xlsx


def _receipt(
    *,
    merchant_name: str,
    total: float,
    category: str,
    currency: str = "USD",
    date: str = "2026-01-15",
    time: str | None = None,
) -> Receipt:
    return Receipt(
        telegram_user_id=1,
        country_code="US",
        merchant_name=merchant_name,
        currency=currency,
        date=date,
        time=time,
        total=total,
        category=category,
        files=ReceiptFiles(original_r2_key="k"),
        created_at=datetime.now(UTC),
    )


def _rows(sheet: object) -> list[tuple[object, ...]]:
    return [tuple(cell.value for cell in row) for row in sheet.iter_rows()]  # type: ignore[attr-defined]


def test_workbook_has_a_summary_sheet_and_one_sheet_per_category() -> None:
    receipts = [
        _receipt(merchant_name="Landlord LLC", total=1200.0, category="room"),
        _receipt(merchant_name="Trader Joes", total=85.32, category="board"),
    ]

    workbook = load_workbook(io.BytesIO(build_report_xlsx(receipts, "2026-01-01", "2026-01-31")))

    assert workbook.sheetnames == ["Summary", "Room (Rent)", "Board (Food)"]


def test_summary_sheet_lists_a_subtotal_per_category_and_a_grand_total() -> None:
    receipts = [
        _receipt(merchant_name="Landlord LLC", total=1200.0, category="room"),
        _receipt(merchant_name="Trader Joes", total=85.32, category="board"),
    ]

    workbook = load_workbook(io.BytesIO(build_report_xlsx(receipts, "2026-01-01", "2026-01-31")))
    rows = _rows(workbook["Summary"])

    assert ("Room (Rent)", "USD", 1200.0) in rows
    assert ("Board (Food)", "USD", 85.32) in rows
    assert ("Grand total", "USD", 1285.32) in rows


def test_room_sheet_only_contains_room_receipts() -> None:
    receipts = [
        _receipt(
            merchant_name="Landlord LLC",
            total=1200.0,
            category="room",
            date="2026-01-05",
            time="09:30",
        ),
        _receipt(merchant_name="Trader Joes", total=85.32, category="board"),
    ]

    workbook = load_workbook(io.BytesIO(build_report_xlsx(receipts, "2026-01-01", "2026-01-31")))
    rows = _rows(workbook["Room (Rent)"])

    assert ("Vendor", "Date", "Time", "Total", "Currency") in rows
    assert ("Landlord LLC", "2026-01-05", "09:30", 1200.0, "USD") in rows
    assert not any("Trader Joes" in row for row in rows)


def test_board_sheet_only_contains_board_receipts() -> None:
    receipts = [
        _receipt(merchant_name="Landlord LLC", total=1200.0, category="room"),
        _receipt(merchant_name="Trader Joes", total=85.32, category="board"),
    ]

    workbook = load_workbook(io.BytesIO(build_report_xlsx(receipts, "2026-01-01", "2026-01-31")))
    rows = _rows(workbook["Board (Food)"])

    assert ("Trader Joes", "2026-01-15", "?", 85.32, "USD") in rows
    assert not any("Landlord LLC" in row for row in rows)


def test_category_sheet_includes_a_subtotal_row() -> None:
    receipts = [
        _receipt(merchant_name="Store A", total=10.0, category="board"),
        _receipt(merchant_name="Store B", total=20.0, category="board"),
    ]

    workbook = load_workbook(io.BytesIO(build_report_xlsx(receipts, "2026-01-01", "2026-01-31")))
    rows = _rows(workbook["Board (Food)"])

    assert (None, None, "Subtotal", 30.0, "USD") in rows


def test_subtotals_are_grouped_by_currency_not_summed_together() -> None:
    receipts = [
        _receipt(merchant_name="US Landlord", total=1000.0, category="room", currency="USD"),
        _receipt(merchant_name="FR Landlord", total=800.0, category="room", currency="EUR"),
    ]

    workbook = load_workbook(io.BytesIO(build_report_xlsx(receipts, "2026-01-01", "2026-01-31")))
    summary_rows = _rows(workbook["Summary"])
    room_rows = _rows(workbook["Room (Rent)"])

    assert ("Room (Rent)", "USD", 1000.0) in summary_rows
    assert ("Room (Rent)", "EUR", 800.0) in summary_rows
    assert ("Grand total", "USD", 1000.0) in summary_rows
    assert ("Grand total", "EUR", 800.0) in summary_rows
    assert (None, None, "Subtotal", 1000.0, "USD") in room_rows
    assert (None, None, "Subtotal", 800.0, "EUR") in room_rows


def test_empty_receipt_list_still_produces_a_valid_workbook() -> None:
    workbook = load_workbook(io.BytesIO(build_report_xlsx([], "2026-01-01", "2026-01-31")))

    assert workbook.sheetnames == ["Summary", "Room (Rent)", "Board (Food)"]
    assert _rows(workbook["Room (Rent)"])[0] == ("Vendor", "Date", "Time", "Total", "Currency")
